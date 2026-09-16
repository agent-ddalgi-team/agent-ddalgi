"""문서 출력기 — D 담당. LLM을 호출하지 않고 최종 ProfileResult(dict)를 MD/DOCX 파일로 저장한다.

경계(contracts/contract.md, day2_addendum §8):
    render_document(profile, output_dir, format) -> 실제로 저장된 파일 Path

규칙(작업기준 05_D §02①, §08)
- 본문은 draft_sections의 순서·문장을 그대로 쓴다. 문장을 새로 쓰거나 fact_id를 다시 매기지 않는다.
- 제목("회사소개서 초안"), 테스트 데이터 표시(is_mock), "내부 검토용" 문구,
  needs_confirmation 확인 질문, sources+evidence 근거 부록을 포함한다.
- 화면(frontend/app.js)과 같은 라벨·같은 본문을 쓴다(contract §6 문서·화면 동일성).
- LLM을 재호출하지 않는다. 지원하지 않는 형식·잘못된 결과·저장 실패는 숨기지 말고 예외로 낸다.
  서버(main.py, C)는 이 예외를 약속된 DOCUMENT_FAILED로 변환한다.
- 저장 위치는 서버가 만든 작업 폴더(output_dir)로 한정하고, 사용자 파일명/경로로 덮어쓰지 않는다.
  기본 파일명은 회사소개서_초안.md / 회사소개서_초안.docx.
- 저장을 마친 뒤에만 실제 경로를 반환한다.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

# 화면(app.js FIELD_LABELS)과 반드시 같은 표기. 확인 질문 라벨에 쓴다.
FIELD_LABELS: dict[str, str] = {
    "company_name": "회사명",
    "company_summary": "회사 개요",
    "business_areas": "사업 분야",
    "products_services": "제품·서비스",
    "technology": "기술",
    "strengths": "강점",
    "customers_markets": "고객·시장",
    "certifications": "인증·승인·특허",
    "history": "연혁",
    "processes": "공정 목록",
    "process_count": "공정 수",
    "capabilities": "대응 범위",
    "lead_time": "납기",
    "other_info": "기타 핵심 정보",
}

TITLE = "회사소개서 초안"
MOCK_BADGE = "테스트 데이터"
INTERNAL_NOTICE = "내부 검토용 · 담당자 확인 전 대외 사용 금지"
DEFAULT_BASENAME = "회사소개서_초안"
SUPPORTED_FORMATS = ("md", "docx")


class DocumentRenderError(Exception):
    """문서를 만들 수 없을 때. 서버는 이를 DOCUMENT_FAILED로 변환한다.

    code: 오류 원인 구분용 짧은 이름. message: 사람이 읽을 설명.
    """

    def __init__(self, code: str, message: str):
        super().__init__(f"[{code}] {message}")
        self.code = code
        self.message = message


# --------------------------------------------------------------------------
# 입력 점검 — 자동으로 고치지 않는다. 문제가 있으면 예외로만 알린다.
# --------------------------------------------------------------------------
def _require(condition: bool, code: str, message: str) -> None:
    if not condition:
        raise DocumentRenderError(code, message)


def _check_profile_shape(profile: Any) -> None:
    _require(isinstance(profile, dict), "bad_profile", "profile은 객체(dict)여야 합니다.")

    sections = profile.get("draft_sections")
    _require(
        isinstance(sections, list) and len(sections) >= 1,
        "bad_profile",
        "draft_sections가 비어 있어 문서 본문을 만들 수 없습니다.",
    )
    for i, section in enumerate(sections):
        _require(isinstance(section, dict), "bad_profile", f"draft_sections[{i}] 형식이 잘못되었습니다.")
        _require(
            isinstance(section.get("title"), str) and section["title"] != "",
            "bad_profile",
            f"draft_sections[{i}].title이 없습니다.",
        )
        paragraphs = section.get("paragraphs")
        _require(
            isinstance(paragraphs, list) and len(paragraphs) >= 1,
            "bad_profile",
            f"draft_sections[{i}].paragraphs가 비어 있습니다.",
        )
        for j, para in enumerate(paragraphs):
            _require(
                isinstance(para, dict) and isinstance(para.get("text"), str) and para["text"] != "",
                "bad_profile",
                f"draft_sections[{i}].paragraphs[{j}].text가 없습니다.",
            )

    _require(isinstance(profile.get("sources"), list), "bad_profile", "sources 목록이 없습니다.")
    _require(isinstance(profile.get("company_info"), dict), "bad_profile", "company_info가 없습니다.")
    # needs_confirmation은 비어 있어도 됨(질문 없음). 존재 시 형식만 확인.
    nc = profile.get("needs_confirmation", [])
    _require(isinstance(nc, list), "bad_profile", "needs_confirmation은 배열이어야 합니다.")


# --------------------------------------------------------------------------
# 공통 데이터 정리 — MD/DOCX가 같은 내용을 쓰도록 한곳에서 만든다.
# --------------------------------------------------------------------------
def _confirmation_lines(profile: dict) -> list[str]:
    lines = []
    for item in profile.get("needs_confirmation", []) or []:
        field = item.get("field", "")
        label = FIELD_LABELS.get(field, field)
        question = item.get("question", "")
        lines.append(f"[{label}] {question}")
    return lines


def _source_name_map(profile: dict) -> dict[str, str]:
    return {s.get("source_id", ""): s.get("file_name", "") for s in profile.get("sources", []) or []}


def _source_rows(profile: dict) -> list[str]:
    return [f"{s.get('file_name', '')} ({s.get('source_id', '')})" for s in profile.get("sources", []) or []]


def _evidence_rows(profile: dict) -> list[str]:
    """company_info의 모든 사실에 달린 evidence를 화면(app.js renderSources)과 같은 형식으로 편다."""
    name_by_id = _source_name_map(profile)
    rows: list[str] = []
    for field in (profile.get("company_info") or {}).values():
        for fact in (field or {}).get("facts", []) or []:
            for ev in fact.get("evidence", []) or []:
                sid = ev.get("source_id", "")
                file_name = name_by_id.get(sid, sid)
                rows.append(f'{file_name} · {ev.get("locator", "")}: "{ev.get("quote", "")}"')
    return rows


# --------------------------------------------------------------------------
# Markdown
# --------------------------------------------------------------------------
def _build_markdown(profile: dict) -> str:
    out: list[str] = []
    title = TITLE + (f" ({MOCK_BADGE})" if profile.get("is_mock") else "")
    out.append(f"# {title}")
    out.append("")
    out.append(f"> {INTERNAL_NOTICE}")
    out.append("")

    out.append("## 소개서 본문")
    out.append("")
    for section in profile["draft_sections"]:
        out.append(f"### {section['title']}")
        for para in section["paragraphs"]:
            out.append(para["text"])  # 문장 그대로. 다시 쓰지 않는다.
        out.append("")

    out.append("## 확인 질문")
    out.append("")
    confirmations = _confirmation_lines(profile)
    if confirmations:
        out.extend(f"- {line}" for line in confirmations)
    else:
        out.append("확인이 필요한 항목이 없습니다.")
    out.append("")

    out.append("## 근거 자료")
    out.append("")
    out.append("### 자료 목록")
    source_rows = _source_rows(profile)
    out.extend(f"- {row}" for row in source_rows) if source_rows else out.append("- 자료 없음")
    out.append("")
    out.append("### 인용 근거")
    evidence_rows = _evidence_rows(profile)
    out.extend(f"- {row}" for row in evidence_rows) if evidence_rows else out.append("- 근거 없음")
    out.append("")

    return "\n".join(out)


# --------------------------------------------------------------------------
# DOCX (python-docx는 docx 형식일 때만 필요하므로 지연 import)
# --------------------------------------------------------------------------
def _set_korean_font(document, font_name: str = "Malgun Gothic") -> None:
    from docx.oxml.ns import qn

    style = document.styles["Normal"]
    style.font.name = font_name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement

        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)


def _build_docx(profile: dict, path: Path) -> None:
    try:
        from docx import Document
    except ModuleNotFoundError as exc:
        raise DocumentRenderError(
            "docx_unavailable",
            "DOCX 출력에는 python-docx가 필요합니다. requirements.txt에 python-docx를 추가하세요.",
        ) from exc

    document = Document()
    _set_korean_font(document)

    title = TITLE + (f" ({MOCK_BADGE})" if profile.get("is_mock") else "")
    document.add_heading(title, level=0)
    note = document.add_paragraph()
    run = note.add_run(INTERNAL_NOTICE)
    run.italic = True

    document.add_heading("소개서 본문", level=1)
    for section in profile["draft_sections"]:
        document.add_heading(section["title"], level=2)
        for para in section["paragraphs"]:
            document.add_paragraph(para["text"])  # 문장 그대로.

    document.add_heading("확인 질문", level=1)
    confirmations = _confirmation_lines(profile)
    if confirmations:
        for line in confirmations:
            document.add_paragraph(line, style="List Bullet")
    else:
        document.add_paragraph("확인이 필요한 항목이 없습니다.")

    document.add_heading("근거 자료", level=1)
    document.add_heading("자료 목록", level=2)
    source_rows = _source_rows(profile)
    if source_rows:
        for row in source_rows:
            document.add_paragraph(row, style="List Bullet")
    else:
        document.add_paragraph("자료 없음")
    document.add_heading("인용 근거", level=2)
    evidence_rows = _evidence_rows(profile)
    if evidence_rows:
        for row in evidence_rows:
            document.add_paragraph(row, style="List Bullet")
    else:
        document.add_paragraph("근거 없음")

    document.save(str(path))


# --------------------------------------------------------------------------
# 공개 함수
# --------------------------------------------------------------------------
def render_document(profile: dict, output_dir, format: str = "md") -> Path:
    """profile을 format(md|docx) 파일로 output_dir에 저장하고, 저장된 실제 경로를 반환한다.

    실패 시 DocumentRenderError를 낸다(성공 값을 임의로 만들지 않는다).
    """
    fmt = str(format).lower()
    _require(fmt in SUPPORTED_FORMATS, "unsupported_format", f"지원하지 않는 형식입니다: {format}")

    _check_profile_shape(profile)

    out_dir = Path(output_dir)
    try:
        out_dir.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise DocumentRenderError("save_failed", f"저장 폴더를 만들 수 없습니다: {out_dir}") from exc

    path = out_dir / f"{DEFAULT_BASENAME}.{fmt}"

    try:
        if fmt == "md":
            path.write_text(_build_markdown(profile), encoding="utf-8")
        else:  # docx
            _build_docx(profile, path)
    except DocumentRenderError:
        raise
    except OSError as exc:
        raise DocumentRenderError("save_failed", f"파일 저장에 실패했습니다: {path}") from exc

    # 저장이 실제로 됐는지 확인한 뒤에만 경로를 돌려준다.
    if not path.exists() or path.stat().st_size == 0:
        raise DocumentRenderError("save_failed", f"파일이 생성되지 않았습니다: {path}")

    return path.resolve()
